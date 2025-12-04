from django.shortcuts import render
from django.http import JsonResponse
from django.db.models import Q, Sum, F, DecimalField
from django.db.models.functions import Coalesce
from django.contrib.auth.decorators import login_required
from mrp.models import RequestItem, PurchaseRequest
import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font, Alignment, PatternFill


@login_required
def request_items_report(request):
    """صفحه اصلی گزارش اقلام خریداری شده"""
    return render(request, 'mrp/purchase/request_items_report.html')


@login_required
def request_items_data(request):
    """API برای دریافت داده‌های جدول با Ajax"""
    
    # پارامترهای DataTables
    draw = int(request.GET.get('draw', 1))
    start = int(request.GET.get('start', 0))
    length = int(request.GET.get('length', 20))
    search_value = request.GET.get('search[value]', '')
    
    # فیلترهای ستون‌ها
    user_filter = request.GET.get('columns[0][search][value]', '')
    item_filter = request.GET.get('columns[1][search][value]', '')
    place_filter = request.GET.get('columns[2][search][value]', '')
    supplier_filter = request.GET.get('columns[3][search][value]', '')
    quantity_filter = request.GET.get('columns[4][search][value]', '')
    price_filter = request.GET.get('columns[5][search][value]', '')
    status_filter = request.GET.get('columns[7][search][value]', '')
    date_filter = request.GET.get('columns[8][search][value]', '')
    
    # Query اصلی
    queryset = RequestItem.objects.select_related(
        'purchase_request__user',
        'item_name',
        'consume_place',
        'supplier_assigned'
    ).all()
    
    # جستجوی عمومی
    if search_value:
        queryset = queryset.filter(
            Q(purchase_request__user__fullName__icontains=search_value) |
            Q(item_name__partName__icontains=search_value) |
            Q(consume_place__assetName__icontains=search_value) |
            Q(supplier_assigned__name__icontains=search_value) |
            Q(description__icontains=search_value)
        )
    
    # فیلترهای ستونی
    if user_filter:
        queryset = queryset.filter(purchase_request__user__fullName__icontains=user_filter)
    if item_filter:
        queryset = queryset.filter(item_name__partName__icontains=item_filter)
    if place_filter:
        queryset = queryset.filter(consume_place__assetName__icontains=place_filter)
    if supplier_filter:
        queryset = queryset.filter(supplier_assigned__name__icontains=supplier_filter)
    if quantity_filter:
        queryset = queryset.filter(quantity=quantity_filter)
    if price_filter:
        queryset = queryset.filter(price=price_filter)
    if status_filter:
        queryset = queryset.filter(purchase_request__status__icontains=status_filter)
    if date_filter:
        queryset = queryset.filter(purchase_request__created_at__icontains=date_filter)
    
    # مرتب‌سازی
    order_column = int(request.GET.get('order[0][column]', 0))
    order_dir = request.GET.get('order[0][dir]', 'desc')
    
    order_columns = {
        0: 'purchase_request__user__fullName',
        1: 'item_name__partName',
        2: 'consume_place__assetName',
        3: 'supplier_assigned__name',
        4: 'quantity',
        5: 'price',
        7: 'purchase_request__status',
        8: 'purchase_request__created_at',
    }
    
    order_by = order_columns.get(order_column, 'purchase_request__created_at')
    if order_dir == 'desc':
        order_by = f'-{order_by}'
    
    queryset = queryset.order_by(order_by)
    
    # محاسبه جمع کل (قبل از pagination)
    total_records = queryset.count()
    total_price = queryset.aggregate(
        total=Sum(F('price') * F('quantity'))
    )['total'] or 0
    
    # Pagination
    paginated_data = queryset[start:start + length]
    
    # آماده‌سازی داده‌ها برای ارسال
    data = []
    for item in paginated_data:
        total_item_price = (item.price or 0) * item.quantity
        
        # وضعیت درخواست
        status_badge = {
            'Pending': 'badge-info',
            'Approved': 'badge-success',
            'Approve2': 'badge-warning',
            'Approve3': 'badge-primary',
            'Approve5': 'badge-secondary',
            'Ordered': 'badge-dark',
            'Purchased': 'badge-success',
            'GuardApproved': 'badge-info',
            'Completed': 'badge-success',
        }.get(item.purchase_request.status, 'badge-secondary')
        
        status_text = dict(item.purchase_request._meta.get_field('status').choices).get(
            item.purchase_request.status, 
            item.purchase_request.status
        )
        
        data.append({
            'user': item.purchase_request.user.fullName,
            'item_name': item.item_name.partName,
            'consume_place': item.consume_place.assetName,
            'supplier': item.supplier_assigned.name if item.supplier_assigned else '<span class="text-muted">تعیین نشده</span>',
            'quantity': item.quantity,
            'price': f'{item.price:,.0f}',
            'total_price': f'{total_item_price:,.0f}',
            'status': f'<span class="badge {status_badge}">{status_text}</span>',
            'date': item.purchase_request.get_dateCreated_jalali().strftime('%Y/%m/%d'),
            'description': item.description or '-',
            'request_id': item.purchase_request.id,
        })
    
    # پاسخ JSON
    response = {
        'draw': draw,
        'recordsTotal': RequestItem.objects.count(),
        'recordsFiltered': total_records,
        'data': data,
        'totalPrice': f'{total_price:,.0f}',
    }
    
    return JsonResponse(response)


@login_required
def export_request_items(request):
    """Export به Excel"""
    
    # دریافت فیلترها از request
    user_filter = request.GET.get('user', '')
    item_filter = request.GET.get('item', '')
    place_filter = request.GET.get('place', '')
    supplier_filter = request.GET.get('supplier', '')
    
    # Query
    queryset = RequestItem.objects.select_related(
        'purchase_request__user',
        'item_name',
        'consume_place',
        'supplier_assigned'
    ).all()
    
    if user_filter:
        queryset = queryset.filter(purchase_request__user__fullName__icontains=user_filter)
    if item_filter:
        queryset = queryset.filter(item_name__partName__icontains=item_filter)
    if place_filter:
        queryset = queryset.filter(consume_place__assetName__icontains=place_filter)
    if supplier_filter:
        queryset = queryset.filter(supplier_assigned__name__icontains=supplier_filter)
    
    # ساخت فایل Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "گزارش خرید"
    
    # استایل‌ها
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    
    # هدر
    headers = ['ردیف', 'سفارش دهنده', 'نام کالا', 'مکان مصرف', 'تامین کننده', 'تعداد', 'قیمت واحد', 'قیمت کل', 'وضعیت', 'تاریخ', 'توضیحات']
    ws.append(headers)
    
    # استایل هدر
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # داده‌ها
    total_price = 0
    for idx, item in enumerate(queryset, 1):
        item_total = (item.price or 0) * item.quantity
        total_price += item_total
        
        status_text = dict(item.purchase_request._meta.get_field('status').choices).get(
            item.purchase_request.status
        )
        
        ws.append([
            idx,
            item.purchase_request.user.fullName,
            item.item_name.partName,
            item.consume_place.assetName,
            item.supplier_assigned.name if item.supplier_assigned else 'تعیین نشده',
            item.quantity,
            item.price or 0,
            item_total,
            status_text,
            item.purchase_request.get_dateCreated_jalali().strftime('%Y/%m/%d'),
            item.description or '-',
        ])
    
    # ردیف جمع
    ws.append(['', '', '', '', '', '', '', total_price, '', '', ''])
    last_row = ws.max_row
    ws[f'A{last_row}'] = 'جمع کل:'
    ws[f'A{last_row}'].font = Font(bold=True)
    ws[f'H{last_row}'].font = Font(bold=True, color="FF0000")
    
    # تنظیم عرض ستون‌ها
    column_widths = [8, 25, 30, 25, 25, 10, 15, 15, 15, 15, 30]
    for idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    
    # ذخیره و ارسال
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=request_items_report.xlsx'
    wb.save(response)
    
    return response